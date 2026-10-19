"""Checks against the committed demo runtime bundle (real archive slice, frozen models)."""

from pathlib import Path

import pandas as pd
import pytest

from moscollector import service as service_module
from moscollector.service import prediction_id

BUNDLE = Path(__file__).resolve().parents[1] / "vercel_runtime"


@pytest.fixture(scope="module")
def service():
    if not (BUNDLE / "artifacts" / "evaluation_report.json").exists():
        pytest.skip("Runtime bundle is not present")
    patch = pytest.MonkeyPatch()
    patch.setattr(service_module, "ARTIFACTS", BUNDLE / "artifacts")
    patch.setattr(service_module, "PROCESSED", BUNDLE / "data" / "processed")
    try:
        yield service_module.AnalyticsService()
    finally:
        patch.undo()


@pytest.mark.parametrize("kind", ["fault", "fire", "access"])
def test_threshold_preview_reproduces_the_released_selection_period(service, kind):
    """The preview replays the scores the active 24-hour-rule bundle was selected on."""
    stored = service.rule[kind]["selection"]
    preview = service.threshold_preview(kind, service.rule[kind]["threshold"])["current"]
    assert preview["threshold"] == service.thresholds()[kind] == service.rule[kind]["threshold"]
    for key in ("precision", "recall", "recall_strict", "alerts_per_day", "false_alerts_per_day"):
        assert preview[key] == pytest.approx(stored[key]), key


def test_archive_queue_issues_at_most_one_warning_per_object_and_kind_a_day(service):
    """Warnings keep their id for 24 hours, so a decision stays with the warning."""
    times = [t for t in service.times if pd.Timestamp("2026-06-14") <= t < pd.Timestamp("2026-06-17")]
    seen = {}
    for t in times:
        for row in service.forecast_rows(t):
            if not row["above_threshold"]:
                assert row["warning_issued_at"] is None
                continue
            issued = pd.Timestamp(row["warning_issued_at"])
            assert issued <= t < issued + pd.Timedelta(hours=24)
            assert row["id"] == prediction_id(row["object_id"], row["kind"], issued)
            assert row["valid_until"] == (issued + pd.Timedelta(hours=24)).isoformat()
            seen.setdefault((row["object_id"], row["kind"]), set()).add(issued)
    assert seen, "the replay window should contain warnings"
    for issued in seen.values():
        ordered = sorted(issued)
        assert all(b - a >= pd.Timedelta(hours=24) for a, b in zip(ordered, ordered[1:], strict=False))


def test_higher_threshold_never_issues_more_warnings(service):
    low = service.threshold_preview("access", 0.2)["proposed"]
    high = service.threshold_preview("access", 0.5)["proposed"]
    assert high["alerts"] <= low["alerts"] and high["warnings_now"] <= low["warnings_now"]


def test_summary_and_equipment_are_consistent_with_forecasts(service):
    forecasts = service.forecast_rows(None)
    summary = service.summary(None)
    assert summary["totals"]["warnings"] == sum(r["above_threshold"] for r in forecasts)
    assert sum(n["warnings"] for n in summary["nodes"]) == summary["totals"]["warnings"]
    assert {s["year"] for s in summary["seasonality"]} >= {2019, 2026}
    equipment = service.equipment(None)
    assert len(equipment["objects"]) == sum(r["kind"] == "fault" for r in forecasts)
    assert all(c["ts"] < forecasts[0]["as_of"] for o in equipment["objects"] for c in o["channels"])


def test_card_context_uses_only_the_past(service):
    detail = service.detail(3215, "access")
    assert detail["calendar"]["object_last_episode"] < detail["as_of"]
    assert {s["id"] for s in detail["verification"]["external_sources"]} == {
        "cameras",
        "planned_works",
        "access_permits",
    }
    assert detail["recommendation_details"][-1]["source"] == "general"


def test_management_reports_render_in_both_formats(service):
    from io import BytesIO

    from openpyxl import load_workbook

    from moscollector.reports import build_report, to_pdf, to_xlsx

    summary = service.summary(None)
    forecasts = service.forecast_rows(None)
    decision = {
        "object_name": "объект Йота",
        "kind": "access",
        "forecast_at": summary["as_of"],
        "action": "false_alarm",
        "action_label": "Ложное срабатывание",
        "reason": "false_signal",
        "comment": "Проверено по камерам",
        "work_order": None,
    }
    report = build_report(summary, forecasts, [decision], [])
    assert report["workload"]["false_alarm_share"] == 1
    book = load_workbook(BytesIO(to_xlsx(report)))
    assert book.sheetnames[:3] == ["Сводка", "По узлам", "Предупреждения"]
    assert book["Предупреждения"].max_row - 1 == summary["totals"]["warnings"]
    pdf = to_pdf(report)
    assert pdf.startswith(b"%PDF") and len(pdf) > 5000


@pytest.mark.parametrize("kind", ["fault", "fire", "access"])
def test_card_confidence_reproduces_the_selection_precision(service, kind):
    threshold = service.thresholds()[kind]
    value = service.confidence(kind, threshold, threshold)
    assert value["warnings"] > 0
    assert value["confirmed"] / value["warnings"] == pytest.approx(
        service.rule[kind]["selection"]["precision"]
    )
    assert value["level"] in ("high", "critical")
    assert value["level_confirmed"] <= value["level_warnings"] <= value["warnings"]


def test_dispatcher_load_shows_the_24_hour_rule_effect(service):
    load = service.summary(None)["load"]
    assert load["after_per_day"] < load["before_per_day"]
    rows = {r["kind"]: r for r in load["rows"]}
    assert not rows["fault"]["changed"] and rows["fault"]["before_per_day"] == rows["fault"]["after_per_day"]
    assert rows["access"]["after_per_day"] < rows["access"]["before_per_day"]


def test_kind_released_with_another_model_is_rescored_over_the_archive(service):
    """Queue, card and new data use the released model, not the archive classifier."""
    from moscollector.model_registry import load_bundle

    rescored = {k for k, m in service.rule.items() if m["source_model"] != f"models/{k}.cbm"}
    assert service.rescored == rescored
    for kind in rescored:
        head = load_bundle(service.rule_version)[kind]
        t = service.resolve_time()
        rows = service.predictions[service.predictions.kind.eq(kind) & service.predictions.as_of.eq(t)]
        frame = rows[["object_id", "as_of"]].merge(service.features, on=["object_id", "as_of"])
        assert len(rows) and rows.probability.to_numpy() == pytest.approx(head.probability(frame))
        assert (rows.threshold == service.rule[kind]["threshold"]).all()
        card = service.detail(3215, kind)
        assert {c["feature"] for c in card["explanation"]} <= set(head.meta["features"])


def test_real_archive_reports_registries_as_not_provided(service):
    detail = service.detail(3215, "access")
    statuses = {s["id"]: s["status"] for s in detail["verification"]["external_sources"]}
    assert set(statuses.values()) == {"not_provided"}


def test_units_report_the_last_journal_state_of_pumps_and_fans(service):
    units = service.units()
    assert units["pumps"] == 180 and units["fans"] == 354
    assert sum(units["states"].values()) == units["pumps"] + units["fans"]
    for obj in units["objects"]:
        assert sum(obj["states"].values()) == obj["pumps"] + obj["fans"]
        assert len(obj["problems"]) == min(8, obj["states"]["problem"])


def test_shift_never_counts_outcomes_after_the_moment(service):
    t = pd.Timestamp("2026-06-15T12:00:00")
    shift = service.shift(t)
    current, resolved = shift["current"]["totals"], shift["resolved"]["totals"]
    assert current["warnings"] == current["correct"] + current["false"] + current["pending"]
    # Results come from a shift that ended a day before the current one began.
    assert pd.Timestamp(shift["resolved"]["begin"]) + pd.Timedelta(hours=36) == pd.Timestamp(
        shift["current"]["begin"]
    )
    assert resolved["pending"] == 0 and resolved["caught"] <= resolved["incident_days"]
    # The same shift seen later can only resolve pending warnings, not change counted ones.
    later = service.shift(t + pd.Timedelta(days=2), None)
    assert later["average"]["shifts"] == shift["average"]["shifts"] == 60
    assert service.shift(t)["current"] == shift["current"]


def test_probability_blend_is_the_weighted_mean_of_calibrated_members(service):
    from moscollector.inference import calibrated, model_input
    from moscollector.model_registry import load_bundle

    for head in load_bundle(service.rule_version).values():
        if not head.blend:
            continue
        rows = service.features[service.features.as_of.eq(service.resolve_time())].head(20)
        expected = sum(
            w * calibrated(m.predict(model_input(rows, f), prediction_type="RawFormulaVal"), c)
            for w, m, f, c in head.members
        )
        assert head.probability(rows) == pytest.approx(expected)
        assert sum(w for w, *_ in head.members) == pytest.approx(1)
        assert head.shap(rows.head(1)).shape == (1, len(head.meta["features"]) + 1)
