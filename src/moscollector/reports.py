"""Management reports (section 8 of the task): XLSX workbook and PDF summary.

Both are built from the same report dictionary, so the two formats never disagree.
"""

from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path

from moscollector.domain import ACTIONS, KIND_LABELS, RISK_LEVELS, WORK_OUTCOMES, WORK_STATUSES

ASSETS = Path(__file__).resolve().parent / "assets"
MONTHS = ["янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"]


def _pct(value):
    return "—" if value is None else f"{value * 100:.1f}%".replace(".", ",")


def _num(value):
    return f"{value:.3f}".replace(".", ",")


def _moment(value: str):
    return str(value)[:16].replace("T", " ")


def build_report(summary: dict, forecasts: list[dict], decisions: list[dict], orders: list[dict]):
    """Plain data for both formats; everything is already computed by the service."""
    actions = {key: sum(d["action"] == key for d in decisions) for key in ACTIONS}
    reviewed = len(decisions)
    return {
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "as_of": summary["as_of"],
        "totals": summary["totals"],
        "kinds": summary["kinds"],
        "nodes": summary["nodes"],
        "quality": summary["quality"],
        "seasonality": summary["seasonality"],
        "warnings": [f for f in forecasts if f["above_threshold"]],
        "decisions": decisions,
        "orders": orders,
        "workload": {
            "decisions": reviewed,
            "actions": actions,
            "false_alarm_share": actions.get("false_alarm", 0) / reviewed if reviewed else None,
            "orders_by_status": {k: sum(o["status"] == k for o in orders) for k in WORK_STATUSES},
            "orders_by_outcome": {k: sum(o["outcome"] == k for o in orders) for k in WORK_OUTCOMES},
        },
    }


def to_xlsx(report: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    book = Workbook()
    bold = Font(bold=True)

    def sheet(title, header, rows, first=False):
        ws = book.active if first else book.create_sheet()
        ws.title = title
        ws.append(header)
        for cell in ws[1]:
            cell.font = bold
        for row in rows:
            ws.append(row)
        for i, column in enumerate(header, 1):
            width = max(
                [len(str(column))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None] or [8]
            )
            ws.column_dimensions[get_column_letter(i)].width = min(60, width + 2)
        ws.freeze_panes = "A2"
        return ws

    totals, workload = report["totals"], report["workload"]
    sheet(
        "Сводка",
        ["Показатель", "Значение"],
        [
            ["Момент прогноза, МСК", _moment(report["as_of"])],
            ["Отчёт сформирован", report["generated_at"]],
            ["Прогнозов (объект × тип)", totals["forecasts"]],
            ["Предупреждений на 24 часа", totals["warnings"]],
            ["Критических", totals["critical"]],
            ["Объектов с предупреждениями", totals["objects_at_risk"]],
            ["Решений диспетчеров в журнале", workload["decisions"]],
            ["Доля ложных срабатываний среди решений", _pct(workload["false_alarm_share"])],
        ]
        + [[f"Решений «{ACTIONS[k]}»", v] for k, v in workload["actions"].items()]
        + [[f"Заявок: {WORK_STATUSES[k]}", v] for k, v in workload["orders_by_status"].items()],
        first=True,
    )
    sheet(
        "По узлам",
        ["Эксплуатационный узел", "Объектов", "Предупреждений", "Критических"]
        + [KIND_LABELS[k] for k in KIND_LABELS],
        [
            [n["name"], n["objects"], n["warnings"], n["critical"]]
            + [n["by_kind"].get(k, 0) for k in KIND_LABELS]
            for n in report["nodes"]
        ],
    )
    sheet(
        "Предупреждения",
        ["Объект", "Тип риска", "Уровень", "Вероятность", "Порог", "Действует до", "Рекомендация"],
        [
            [
                f["object_name"],
                f["kind_label"],
                RISK_LEVELS.get(f["risk"], f["risk"]),
                round(f["probability"], 4),
                round(min(f["threshold"], 1), 4),
                _moment(f["valid_until"]),
                f["recommendation"],
            ]
            for f in report["warnings"]
        ],
    )
    sheet(
        "Решения",
        ["Объект", "Тип риска", "Момент прогноза", "Действие", "Причина", "Комментарий", "Статус работ"],
        [
            [
                d["object_name"],
                KIND_LABELS.get(d["kind"], d["kind"]),
                _moment(d["forecast_at"]),
                d["action_label"],
                d.get("reason_label", d["reason"]),
                d["comment"],
                (d.get("work_order") or {}).get("status_label", "—"),
            ]
            for d in report["decisions"]
        ],
    )
    sheet(
        "Заявки",
        ["Номер", "Объект", "Тип риска", "Приоритет", "Статус", "Итог", "Создана"],
        [
            [
                o["number"],
                o["object_name"],
                KIND_LABELS.get(o["kind"], o["kind"]),
                o["priority_label"],
                o["status_label"],
                o["outcome_label"] or "—",
                _moment(o["created_at"]),
            ]
            for o in report["orders"]
        ],
    )
    sheet(
        "Качество модели",
        ["Тип риска", "Precision", "Recall", "F1", "F1 базы", "Эпизодов в тесте", "Предупреждения"],
        [
            [
                KIND_LABELS[k],
                round(q["precision"], 3),
                round(q["recall"], 3),
                round(q["f1"], 3),
                round(q["baseline_f1"], 3),
                q["eligible_episodes"],
                "включены" if q["enabled"] else "отключены",
            ]
            for k, q in report["quality"].items()
        ],
    )
    sheet(
        "Сезонность",
        ["Год", "Месяц", "Тревожных сообщений", "Примечание"],
        [
            [
                s["year"],
                MONTHS[s["month"] - 1],
                s["alarms"],
                "миграция мониторинга, исключён" if s["quarantined"] else "",
            ]
            for s in report["seasonality"]
        ],
    )
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def to_pdf(report: dict) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.add_font("PTSans", "", str(ASSETS / "PTSans-Regular.ttf"))
    pdf.add_font("PTSans", "B", str(ASSETS / "PTSans-Bold.ttf"))
    pdf.set_auto_page_break(True, margin=14)
    pdf.set_margins(14, 14, 14)
    pdf.add_page()

    def heading(text, size=12):
        pdf.ln(3)
        pdf.set_font("PTSans", "B", size)
        pdf.cell(0, 7, text, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("PTSans", "", 9.5)

    def table(header, rows, widths):
        pdf.set_font("PTSans", "B", 8.5)
        with pdf.table(col_widths=widths, text_align="LEFT", line_height=5) as grid:
            row = grid.row()
            for value in header:
                row.cell(str(value))
            pdf.set_font("PTSans", "", 8.5)
            for values in rows:
                row = grid.row()
                for value in values:
                    row.cell(str(value))
        pdf.set_font("PTSans", "", 9.5)

    pdf.set_font("PTSans", "B", 16)
    pdf.cell(0, 9, "Отчёт руководителю подразделения", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("PTSans", "", 9.5)
    pdf.cell(
        0,
        5,
        f"Прогноз на {_moment(report['as_of'])} МСК, горизонт 24 часа · сформирован {report['generated_at']}",
        new_x="LMARGIN",
        new_y="NEXT",
    )
    totals, workload = report["totals"], report["workload"]
    heading("Итоги")
    table(
        ["Предупреждений", "Критических", "Объектов с риском", "Решений", "Ложные срабатывания"],
        [
            [
                totals["warnings"],
                totals["critical"],
                totals["objects_at_risk"],
                workload["decisions"],
                _pct(workload["false_alarm_share"]),
            ]
        ],
        (1, 1, 1, 1, 1),
    )
    heading("По типам риска")
    table(
        ["Тип", "Предупреждений", "Критических", "Порог"],
        [
            [
                k["label"],
                k["warnings"],
                k["critical"],
                "отключено" if k["threshold"] > 1 else _pct(k["threshold"]),
            ]
            for k in report["kinds"]
        ],
        (2, 1, 1, 1),
    )
    heading("Эксплуатационные узлы с предупреждениями")
    nodes = [n for n in report["nodes"] if n["warnings"]] or report["nodes"][:5]
    table(
        ["Узел", "Объектов", "Предупреждений", "Критических"],
        [[n["name"], n["objects"], n["warnings"], n["critical"]] for n in nodes[:16]],
        (3, 1, 1, 1),
    )
    heading("Отработка предупреждений")
    table(
        ["Действие", "Решений"],
        [[ACTIONS[k], v] for k, v in workload["actions"].items()]
        + [[f"Заявки: {WORK_STATUSES[k]}", v] for k, v in workload["orders_by_status"].items() if v],
        (3, 1),
    )
    heading("Качество модели на отложенном тесте (июнь 2026)")
    table(
        ["Тип", "Precision", "Recall", "F1", "F1 базы"],
        [
            [KIND_LABELS[k], _pct(q["precision"]), _pct(q["recall"]), _num(q["f1"]), _num(q["baseline_f1"])]
            if q["enabled"]
            else [KIND_LABELS[k], "предупреждения отключены", "—", "—", "—"]
            for k, q in report["quality"].items()
        ],
        (2, 1, 1, 1, 1),
    )
    pdf.ln(2)
    pdf.set_font("PTSans", "", 8)
    pdf.multi_cell(
        0,
        4,
        "Метрики событийные: одно предупреждение сопоставляется с одним эпизодом журнала датчиков. "
        "Реальные инциденты в данных не размечены; прогноз относится к зарегистрированным эпизодам.",
    )
    return bytes(pdf.output())
